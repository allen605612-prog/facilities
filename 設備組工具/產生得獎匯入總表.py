# -*- coding: utf-8 -*-
"""
競賽得獎名單 → 校務系統匯入總表 產生器

用法（預設路徑已設好，直接跑即可）：
    uvx --with openpyxl python 產生得獎匯入總表.py
    uvx --with openpyxl python 產生得獎匯入總表.py 輸入.xlsx 輸出.xlsx

輸入檔欄位（第一列為標題，順序不拘、以標題名稱對應）：
    屆別 科別 名次 題目 班級 學號 座號 姓名 指導老師 備註 比賽日期 證書日期 證書文號
輸出檔為 16 欄匯入格式，欄位型別（文字/數字）比照校務系統範本。
"""
import sys, os, re, shutil, datetime
import openpyxl
from openpyxl.styles import Font, Alignment

# ─────────── 可調參數 ───────────
IN_DEFAULT  = r'D:\D\114設備組\國中科展得獎名單_整理.xlsx'
OUT_DEFAULT = r'D:\D\114設備組\65屆縣科展得獎名單總表.xlsx'

競賽性質 = '全縣'
競賽類別 = '資訊類'      # 系統規定的固定值，勿改
競賽組別 = '個人賽'      # 系統規定的固定值，勿改
提報人   = 'A130'
主辦單位 = '雲林縣政府'
項目樣板 = '雲林縣第{屆}屆公私立國中小科學展覽競賽榮獲國中組{科別}{名次}'

# 班級三碼代碼 = (年級+6)*100 + 班序。班序依名條上的班級順序。
班序預設 = ['正', '心', '誠', '意', '修', '身', '齊', '音甲', '音乙', '美甲', '美乙']
# 某些年級若無「齊」班，序號會往前遞補，在此依「年級」覆寫，例：
# 班序覆寫 = {3: ['正','心','誠','意','修','身','音甲','音乙','美甲','美乙']}
班序覆寫 = {}

標題 = ['*學年', '*學期', '*競賽項目', '*競賽時間', '*競賽性質', '*競賽類別', '*競賽組別',
        '*提報人', '*學號', '*班級', '*座號', '名次', '*說明', '*證書日期', '*證書字號', '*主辦單位']
文字欄 = {1, 2, 5, 8, 9}          # 需存成文字（@）的欄，其餘數字欄存成數值
# ────────────────────────────────

中文數 = {'一':1,'二':2,'三':3,'四':4,'五':5,'六':6,'七':7,'八':8,'九':9,'十':10}
warns = []


def 取名次(s):
    m = re.match(r'^第([一二三四五六七八九十])名$', str(s).strip())
    return 中文數[m.group(1)] if m else None


def 取屆別(s):
    m = re.search(r'(\d+)', str(s))
    return int(m.group(1)) if m else None


def 取班級代碼(s, row):
    m = re.match(r'^(\d)\s*年\s*(.+?)\s*班$', str(s).strip())
    if not m:
        warns.append(f'第{row}列 班級「{s}」格式無法解析（預期如「2年美甲班」）')
        return None
    年級, 班名 = int(m.group(1)), m.group(2)
    序 = 班序覆寫.get(年級, 班序預設)
    if 班名 not in 序:
        warns.append(f'第{row}列 班級「{s}」的「{班名}」不在班序表中，請確認名條或設定「班序覆寫」')
        return None
    return (年級 + 6) * 100 + 序.index(班名) + 1


def 取民國年(日期, row, 欄名):
    s = str(日期).strip()
    if not re.match(r'^\d{7}$', s):
        warns.append(f'第{row}列 {欄名}「{日期}」不是 7 碼民國日期（如 1150425）')
        return None
    return int(s[:3])


def main():
    inp = sys.argv[1] if len(sys.argv) > 1 else IN_DEFAULT
    out = sys.argv[2] if len(sys.argv) > 2 else OUT_DEFAULT
    if not os.path.exists(inp):
        sys.exit(f'找不到輸入檔：{inp}')

    ws = openpyxl.load_workbook(inp, data_only=True).worksheets[0]
    rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
    head = [str(c).strip() if c is not None else '' for c in rows[0]]
    idx = {name: head.index(name) for name in head}
    need = ['屆別', '科別', '名次', '班級', '學號', '座號', '姓名', '比賽日期', '證書日期', '證書文號']
    missing = [n for n in need if n not in idx]
    if missing:
        sys.exit('輸入檔缺少欄位：' + '、'.join(missing))

    out_rows = []
    for i, r in enumerate(rows[1:], 2):
        g = lambda n: r[idx[n]] if idx.get(n) is not None and idx[n] < len(r) else None
        姓名 = g('姓名')
        if not g('學號'):
            warns.append(f'第{i}列（{姓名}）沒有學號，已略過')
            continue
        屆 = 取屆別(g('屆別'))
        名次數 = 取名次(g('名次'))
        代碼 = 取班級代碼(g('班級'), i)
        比賽 = str(g('比賽日期')).strip()
        證日 = str(g('證書日期')).strip()
        民國年 = 取民國年(比賽, i, '比賽日期')
        取民國年(證日, i, '證書日期')
        if 名次數 is None:
            warns.append(f'第{i}列（{姓名}）名次「{g("名次")}」無法轉成數字')
        if 民國年 and 屆 and 民國年 != 屆 + 49:
            warns.append(f'第{i}列（{姓名}）第{屆}屆的比賽日期應為民國{屆+49}年，實際是{民國年}年')
        學年 = (民國年 - 1) if 民國年 else None      # 科展在下學期，民國年 -1 = 學年度
        out_rows.append({
            'sort': (屆 or 0, i),
            'vals': [
                str(學年) if 學年 else '', '2',
                項目樣板.format(屆=屆, 科別=str(g('科別')).strip(), 名次=str(g('名次')).strip()),
                int(比賽) if 比賽.isdigit() else 比賽,
                競賽性質, 競賽類別, 競賽組別, 提報人,
                str(int(g('學號'))) if isinstance(g('學號'), float) else str(g('學號')).strip(),
                代碼, int(g('座號')), 名次數, str(g('名次')).strip(),
                int(證日) if 證日.isdigit() else 證日,
                str(g('證書文號')).strip(), 主辦單位,
            ],
            '姓名': 姓名,
        })

    out_rows.sort(key=lambda x: x['sort'])           # 依屆別排序，同屆維持輸入順序

    if os.path.exists(out):
        try:                                          # 先確認沒被 Excel 開著，避免備份完才失敗
            open(out, 'r+b').close()
        except PermissionError:
            sys.exit(f'輸出檔正被其他程式開啟（通常是 Excel），請先關閉再執行：{out}')
        bak = f'{os.path.splitext(out)[0]}_備份_{datetime.datetime.now():%Y%m%d_%H%M%S}.xlsx'
        shutil.copyfile(out, bak)
        print(f'已備份原檔 → {bak}')

    wb = openpyxl.Workbook()
    o = wb.active
    o.title = '工作表1'
    o.append(標題)
    for c in o[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal='center')
    for rec in out_rows:
        o.append(rec['vals'])
    for row in o.iter_rows(min_row=2):
        for c in row:
            if c.column in 文字欄:
                c.number_format = '@'
                c.value = str(c.value)
    for col, w in {'A': 9, 'B': 7, 'C': 86.875, 'D': 11, 'E': 10, 'F': 10, 'G': 10, 'H': 9,
                   'I': 10, 'J': 8, 'K': 8, 'L': 7, 'M': 9, 'N': 11, 'O': 26.625, 'P': 12}.items():
        o.column_dimensions[col].width = w
    o.freeze_panes = 'A2'
    wb.create_sheet('工作表2')
    wb.save(out)

    print(f'完成：{len(out_rows)} 列 → {out}')
    for 屆 in sorted({r["sort"][0] for r in out_rows}):
        n = sum(1 for r in out_rows if r['sort'][0] == 屆)
        print(f'   第{屆}屆 {n} 人次')
    if warns:
        print(f'\n⚠ 需要確認 {len(warns)} 項：')
        for w in warns:
            print('  -', w)
    else:
        print('\n✓ 所有欄位轉換正常，無警告')


if __name__ == '__main__':
    main()
